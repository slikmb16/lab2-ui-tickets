"""
Веб-приложение «Генератор билетов» (Flask).

Запуск:
    python main.py
"""
from __future__ import annotations

import io
import importlib.util
import logging
import os
import sys
import traceback
import unittest
from typing import Any, Dict, List, Optional

from flask import Flask, jsonify, render_template, request

from . import logic, parsing, storage

logger = logging.getLogger("lab2.web")

STUDENTS_FILE = "students.xlsx"
TICKETS_FILE = "tickets.docx"
RESULTS_FILE = "results.xlsx"

TEST_LABELS = {
    "test_parses_valid_tickets": "Парсинг валидных билетов",
    "test_skips_ticket_with_fewer_than_3_questions": "Билет с < 3 вопросами пропускается",
    "test_unrecognized_header_does_not_crash_parsing": "Нераспознанный заголовок не роняет парсер",
    "test_extra_questions_truncated_to_three": "Лишние вопросы обрезаются до трёх",
    "test_parses_students_per_group": "Студенты читаются по группам (листам)",
    "test_empty_group_returns_empty_list_no_exception": "Пустая группа — пустой список без исключения",
    "test_new_student_gets_random_ticket_marked_not_repeat": "Новому студенту — случайный билет, повтор = нет",
    "test_repeated_student_gets_same_ticket_marked_repeat": "Повторный выбор — тот же билет",
    "test_first_recorded_ticket_wins_even_after_many_repeats": "Первая запись в журнале побеждает",
    "test_old_rows_are_not_modified_on_append": "Старые строки журнала не меняются",
    "test_different_groups_same_name_are_independent": "Одноимённые студенты из разных групп независимы",
    "test_creates_file_with_correct_headers_if_missing": "Создание results.xlsx с заголовками",
    "test_does_not_overwrite_existing_file": "Существующий журнал не перезаписывается",
    "test_missing_input_files_return_clear_fatal_status": "Отсутствующие входные файлы дают понятное сообщение",
    "test_locked_results_file_raises_friendly_error": "Заблокированный журнал обрабатывается без падения",
}

# Подсказки отображаются на вкладке «Тесты»: на защите сразу видно,
# какое требование лабораторной покрывает каждая автоматическая проверка.
TEST_METADATA = {
    "test_parses_valid_tickets": ("Читает два корректно оформленных билета из Word.", "Приёмка: билет и 3 вопроса из файла.", "Найдены оба номера и по три вопроса."),
    "test_skips_ticket_with_fewer_than_3_questions": ("Передаёт билет с одним вопросом и один корректный билет.", "Обязательное: некорректные билеты пропускаются.", "Неполный билет не попадает в список, корректный остаётся."),
    "test_unrecognized_header_does_not_crash_parsing": ("Добавляет строку с неверным заголовком перед правильным билетом.", "Обязательное: неверный заголовок не роняет приложение.", "Парсер игнорирует мусор и читает корректный билет."),
    "test_extra_questions_truncated_to_three": ("Передаёт билет с четырьмя вопросами.", "Приёмка: в результате отображаются три вопроса.", "Для билета сохранены первые 3 вопроса."),
    "test_parses_students_per_group": ("Читает несколько листов Excel со студентами.", "Приёмка: список студентов соответствует выбранному листу.", "Группы и Фамилия Имя прочитаны без смешивания."),
    "test_empty_group_returns_empty_list_no_exception": ("Читает лист только с заголовками.", "Обязательное: пустая группа.", "Возвращается пустой список без ошибки."),
    "test_new_student_gets_random_ticket_marked_not_repeat": ("Назначает билет студенту, которого нет в журнале.", "Логика назначения: новая выдача.", "Выбран существующий билет, повтор = нет."),
    "test_repeated_student_gets_same_ticket_marked_repeat": ("Выдаёт билет студенту дважды.", "Приёмка: повтор возвращает тот же номер.", "Второй номер совпадает с первым, повтор = да."),
    "test_first_recorded_ticket_wins_even_after_many_repeats": ("Повторяет выдачу одному студенту несколько раз.", "Приёмка: полная история повторов.", "Все строки имеют первый номер; первая — нет, остальные — да."),
    "test_old_rows_are_not_modified_on_append": ("Добавляет две записи в журнал.", "Приёмка: старые строки не тронуты.", "Первая запись сохранена, новая добавлена в конец."),
    "test_different_groups_same_name_are_independent": ("Использует одинаковое имя в двух группах.", "Логика: студент определяется по группе и ФИО.", "Во второй группе это новая, а не повторная выдача."),
    "test_creates_file_with_correct_headers_if_missing": ("Запускает работу без results.xlsx.", "Формат results.xlsx.", "Создан журнал с шестью требуемыми колонками."),
    "test_does_not_overwrite_existing_file": ("Создаёт запись, затем повторно вызывает создание журнала.", "Приёмка: история не перезаписывается.", "Существующая строка остаётся на месте."),
    "test_missing_input_files_return_clear_fatal_status": ("Запускает веб-приложение в пустой папке.", "Обязательное: валидация источников.", "Статус содержит имена отсутствующих файлов и приложение отвечает без падения."),
    "test_locked_results_file_raises_friendly_error": ("Имитирует отказ Excel в записи results.xlsx.", "Обязательное: блокировка файла.", "Возвращается понятная ошибка ResultsFileLocked, без аварийного завершения."),
}


def create_app(base_dir: str = ".") -> Flask:
    root = os.path.abspath(base_dir)
    app = Flask(
        __name__,
        template_folder=os.path.join(root, "templates"),
        static_folder=os.path.join(root, "static"),
    )
    app.config["TEMPLATES_AUTO_RELOAD"] = True
    app.jinja_env.auto_reload = True
    app.config["SEND_FILE_MAX_AGE_DEFAULT"] = 0
    # По ТЗ входные файлы лежат рядом с приложением. Не используем скрытые
    # резервные копии: отсутствие файла должно быть сразу заметно пользователю.
    app.config["STUDENTS_PATH"] = os.path.join(root, STUDENTS_FILE)
    app.config["TICKETS_PATH"] = os.path.join(root, TICKETS_FILE)
    app.config["RESULTS_PATH"] = os.path.join(root, RESULTS_FILE)

    state: Dict[str, Any] = {
        "groups": {},
        "tickets": {},
        "fatal": None,
        "status": "",
    }

    def load_sources() -> None:
        # Источники могут быть переименованы или возвращены, пока сервер
        # работает. Сбрасываем старое состояние перед каждой проверкой.
        state["groups"] = {}
        state["tickets"] = {}
        state["fatal"] = None
        state["status"] = ""
        missing = []
        if not os.path.exists(app.config["STUDENTS_PATH"]):
            missing.append(STUDENTS_FILE)
        if not os.path.exists(app.config["TICKETS_PATH"]):
            missing.append(TICKETS_FILE)
        if missing:
            state["fatal"] = (
                "Не найдены необходимые файлы: "
                + ", ".join(missing)
                + f". Поместите их в папку проекта: {root}."
            )
            return

        try:
            state["groups"] = parsing.load_students(app.config["STUDENTS_PATH"])
        except Exception as exc:
            state["fatal"] = f"Не удалось прочитать {STUDENTS_FILE}: {exc}"
            return

        try:
            state["tickets"] = parsing.load_tickets(app.config["TICKETS_PATH"])
        except Exception as exc:
            state["fatal"] = f"Не удалось прочитать {TICKETS_FILE}: {exc}"
            return

        if not state["tickets"]:
            state["fatal"] = f"В файле {TICKETS_FILE} не найдено ни одного корректного билета."
            return

        try:
            storage.ensure_results_file(app.config["RESULTS_PATH"])
        except Exception as exc:
            state["fatal"] = f"Не удалось создать/открыть {RESULTS_FILE}: {exc}"
            return

        state["status"] = (
            f"Загружено групп: {len(state['groups'])}, "
            f"билетов: {len(state['tickets'])}"
        )

    load_sources()

    def _walk_suite(suite: unittest.TestSuite) -> List[unittest.TestCase]:
        tests: List[unittest.TestCase] = []
        for item in suite:
            if isinstance(item, unittest.TestSuite):
                tests.extend(_walk_suite(item))
            else:
                tests.append(item)
        return tests

    def _load_lab2_suite() -> unittest.TestSuite:
        """Загружает tests/test_lab2.py независимо от способа запуска Flask."""
        local_test = os.path.join(root, "tests", "test_lab2.py")
        bundled_test = os.path.join(
            os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
            "tests",
            "test_lab2.py",
        )
        test_path = local_test if os.path.isfile(local_test) else bundled_test
        if not os.path.isfile(test_path):
            raise FileNotFoundError("Не найден файл tests/test_lab2.py")

        spec = importlib.util.spec_from_file_location("lab2_web_tests", test_path)
        if spec is None or spec.loader is None:
            raise ImportError("Не удалось загрузить tests/test_lab2.py")
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        return unittest.defaultTestLoader.loadTestsFromModule(module)

    def _catalog() -> List[Dict[str, str]]:
        suite = _load_lab2_suite()
        catalog = []
        for test in _walk_suite(suite):
            method = test._testMethodName
            catalog.append(
                {
                    "id": test.id(),
                    "name": method,
                    "class_name": test.__class__.__name__,
                    "label": TEST_LABELS.get(method, method),
                    "description": TEST_METADATA.get(method, ("Проверка логики приложения.", "Требование лабораторной.", "Тест должен завершиться успешно."))[0],
                    "criterion": TEST_METADATA.get(method, ("", "Требование лабораторной.", ""))[1],
                    "expected": TEST_METADATA.get(method, ("", "", "Тест должен завершиться успешно."))[2],
                }
            )
        return catalog

    def _run_unittest(test_id: Optional[str] = None) -> Dict[str, Any]:
        project_root = root
        if project_root not in sys.path:
            sys.path.insert(0, project_root)

        suite = _load_lab2_suite()
        if test_id:
            selected = [test for test in _walk_suite(suite) if test.id() == test_id]
            if not selected:
                raise ValueError("Выбранный тест не найден")
            suite = unittest.TestSuite(selected)

        # TextTestRunner очищает TestSuite после прогона, поэтому список
        # нужен заранее для корректного статуса каждой кнопки.
        tests_to_report = _walk_suite(suite)

        stream = io.StringIO()
        runner = unittest.TextTestRunner(stream=stream, verbosity=2)
        result = runner.run(suite)

        details: List[Dict[str, Any]] = []

        for test in tests_to_report:
            method = test._testMethodName
            entry = {
                "id": test.id(),
                "name": method,
                "label": TEST_LABELS.get(method, method),
                "status": "passed",
                "message": "",
            }
            # TextTestRunner doesn't keep per-test objects easily; match by id()
            for failed_test, err in result.failures + result.errors:
                if failed_test.id() == test.id():
                    entry["status"] = "failed"
                    entry["message"] = err
                    break
            for skipped_test, reason in getattr(result, "skipped", []):
                if skipped_test.id() == test.id():
                    entry["status"] = "skipped"
                    entry["message"] = reason
                    break
            details.append(entry)

        return {
            "ok": result.wasSuccessful(),
            "run": result.testsRun,
            "failures": len(result.failures),
            "errors": len(result.errors),
            "details": details,
            "log": stream.getvalue(),
        }

    @app.get("/")
    def index():
        return render_template("index.html")

    @app.get("/api/status")
    def api_status():
        # Это позволяет увидеть проблему сразу после переименования файла,
        # без остановки и повторного запуска сервера.
        load_sources()
        journal_count = 0
        path = app.config["RESULTS_PATH"]
        if os.path.exists(path) and not state["fatal"]:
            try:
                from openpyxl import load_workbook

                wb = load_workbook(path, read_only=True, data_only=True)
                journal_count = max(wb.active.max_row - 1, 0)
                wb.close()
            except Exception:
                journal_count = 0
        return jsonify(
            {
                "fatal": state["fatal"],
                "status": state["status"],
                "groups": list(state["groups"].keys()),
                "tickets": len(state["tickets"]),
                "journal": journal_count,
            }
        )

    @app.get("/api/groups/<group>/students")
    def api_students(group: str):
        if state["fatal"]:
            return jsonify({"error": state["fatal"]}), 400
        students = state["groups"].get(group)
        if students is None:
            return jsonify({"error": f"Группа «{group}» не найдена"}), 404
        return jsonify(
            {
                "group": group,
                "students": [s.full_name for s in students],
                "empty": len(students) == 0,
            }
        )

    @app.post("/api/generate")
    def api_generate():
        if state["fatal"]:
            return jsonify({"error": state["fatal"]}), 400
        payload = request.get_json(silent=True) or {}
        group = (payload.get("group") or "").strip()
        full_name = (payload.get("student") or "").strip()
        if not group or not full_name:
            return jsonify({"error": "Выберите группу и студента"}), 400

        students = state["groups"].get(group)
        if students is None:
            return jsonify({"error": f"Группа «{group}» не найдена"}), 404
        if not any(s.full_name == full_name for s in students):
            return jsonify({"error": "Студент не найден в выбранной группе"}), 404

        # Берём значения из Excel, а не разбираем отображаемое имя строкой.
        # Так фамилия/имя в results.xlsx всегда совпадают с исходным файлом.
        student = next(s for s in students if s.full_name == full_name)
        results_path = app.config["RESULTS_PATH"]

        try:
            ticket_number, repeat = logic.assign_ticket(
                results_path, group, student.surname, student.name, state["tickets"]
            )
        except storage.ResultsFileLocked as exc:
            return jsonify({"error": str(exc)}), 409
        except Exception as exc:
            logger.exception("assign_ticket failed")
            return jsonify({"error": f"Не удалось назначить билет: {exc}"}), 500

        try:
            storage.append_result(
                results_path, group, student.surname, student.name, ticket_number, repeat
            )
        except storage.ResultsFileLocked as exc:
            return jsonify({"error": str(exc)}), 409
        except Exception as exc:
            logger.exception("append_result failed")
            return jsonify({"error": f"Не удалось сохранить результат: {exc}"}), 500

        questions = state["tickets"][ticket_number]
        return jsonify(
            {
                "ticket_number": ticket_number,
                "questions": questions,
                "repeat": repeat,
                "group": group,
                "student": full_name,
            }
        )

    @app.get("/api/results")
    def api_results():
        path = app.config["RESULTS_PATH"]
        if not os.path.exists(path):
            return jsonify({"rows": []})
        from openpyxl import load_workbook

        try:
            wb = load_workbook(path, read_only=True, data_only=True)
            rows = []
            for row in wb.active.iter_rows(min_row=2, values_only=True):
                if not row or all(c is None for c in row):
                    continue
                rows.append(
                    {
                        "group": row[0],
                        "surname": row[1],
                        "name": row[2],
                        "ticket": row[3],
                        "datetime": str(row[4]) if len(row) > 4 else "",
                        "repeat": row[5] if len(row) > 5 else "",
                    }
                )
        except (PermissionError, OSError):
            return jsonify({"error": "Файл results.xlsx открыт в Excel. Закройте его и обновите журнал."}), 409
        finally:
            if "wb" in locals():
                wb.close()
        return jsonify({"rows": rows})

    @app.get("/api/tests")
    def api_tests():
        return jsonify({"tests": _catalog()})

    @app.post("/api/tests/run")
    def api_tests_run():
        payload = request.get_json(silent=True) or {}
        test_id = payload.get("id")
        try:
            return jsonify(_run_unittest(test_id))
        except Exception as exc:
            logger.exception("test run failed")
            return jsonify({"error": str(exc), "trace": traceback.format_exc()}), 500

    return app


def run(base_dir: str = ".", host: str = "127.0.0.1", port: int = 5000) -> None:
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
    )
    app = create_app(base_dir=base_dir)
    print(f"\n  Генератор билетов: http://{host}:{port}\n")
    app.run(host=host, port=port, debug=False)
