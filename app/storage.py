"""
Работа с журналом результатов results.xlsx:
создание файла, поиск предыдущего билета студента, дозапись строки.
"""
import logging
import os
import time
from datetime import datetime
from typing import Optional

from openpyxl import Workbook, load_workbook

logger = logging.getLogger("lab2.storage")

HEADERS = ["Группа", "Фамилия", "Имя", "Номер билета", "Дата и время", "Повтор"]

DATE_FORMAT = "%Y-%m-%d %H:%M:%S"


class ResultsFileLocked(Exception):
    """results.xlsx открыт другим приложением и недоступен для записи."""


def ensure_results_file(path: str) -> None:
    """Создаёт results.xlsx с заголовком, если файл отсутствует."""
    if os.path.exists(path):
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "Результаты"
    ws.append(HEADERS)
    wb.save(path)
    logger.info("Создан новый журнал результатов: %s", path)


def find_first_ticket(path: str, group: str, surname: str, name: str) -> Optional[int]:
    """
    Возвращает номер билета из самой первой (по порядку строк, т.е.
    хронологически самой ранней) записи по данному студенту, либо None,
    если студент ещё не встречался в журнале.
    """
    if not os.path.exists(path):
        return None

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    result = None
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row is None:
            continue
        row_group, row_surname, row_name, row_ticket = (row + (None,) * 4)[:4]
        if (
            str(row_group).strip() == group
            and str(row_surname).strip() == surname
            and str(row_name).strip() == name
        ):
            result = int(row_ticket)
            break  # первая по порядку запись — файл только дописывается, порядок хронологический
    wb.close()
    return result


def append_result(
    path: str,
    group: str,
    surname: str,
    name: str,
    ticket_number: int,
    repeat: bool,
    retries: int = 3,
    retry_delay_sec: float = 1.0,
) -> None:
    """
    Дописывает строку в конец results.xlsx.
    Существующие строки не изменяются.
    При недоступности файла (открыт в Excel) — несколько попыток,
    затем ResultsFileLocked.
    """
    ensure_results_file(path)

    row = [
        group,
        surname,
        name,
        ticket_number,
        datetime.now().strftime(DATE_FORMAT),
        "да" if repeat else "нет",
    ]

    last_error: Optional[Exception] = None
    for attempt in range(1, retries + 1):
        try:
            wb = load_workbook(path)
            ws = wb.active
            ws.append(row)
            wb.save(path)
            logger.info("Строка добавлена в журнал: %s", row)
            return
        except PermissionError as exc:
            last_error = exc
            logger.warning(
                "results.xlsx недоступен для записи (попытка %d/%d): %s",
                attempt, retries, exc,
            )
            if attempt < retries:
                time.sleep(retry_delay_sec)

    raise ResultsFileLocked(
        "Файл results.xlsx открыт в другой программе (например, Excel). "
        "Закройте файл и попробуйте снова."
    ) from last_error
