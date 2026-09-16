"""Создаёт тестовые students.xlsx и tickets.docx в data/ и в корне проекта."""
import os
import shutil

from docx import Document
from openpyxl import Workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, "data")
os.makedirs(DATA, exist_ok=True)


def write_students(path):
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("ИС-21")
    ws.append(["Фамилия", "Имя"])
    ws.append(["Иванов", "Иван"])
    ws.append(["Петрова", "Анна"])
    ws.append(["Козлов", "Дмитрий"])

    ws2 = wb.create_sheet("ПИ-22")
    ws2.append(["Фамилия", "Имя"])
    ws2.append(["Сидоров", "Семён"])
    ws2.append(["Орлова", "Мария"])

    ws3 = wb.create_sheet("АРХ-20")
    ws3.append(["Фамилия", "Имя"])

    wb.save(path)


def write_tickets(path):
    doc = Document()
    tickets = [
        (1, [
            "Что такое алгоритм и каковы его свойства?",
            "Чем стек отличается от очереди?",
            "Что такое рекурсия? Приведите пример.",
        ]),
        (2, [
            "Что описывает асимптотическая сложность?",
            "Как работает бинарный поиск?",
            "Для чего нужна хеш-таблица?",
        ]),
        (3, [
            "В чём разница между процессом и потоком?",
            "Что такое взаимная блокировка (deadlock)?",
            "Зачем нужен планировщик ОС?",
        ]),
        (4, [
            "Этот билет намеренно битый — только два вопроса.",
            "Парсер должен его пропустить и продолжить работу.",
        ]),
    ]
    for number, questions in tickets:
        doc.add_paragraph(f"Билет {number}")
        for i, q in enumerate(questions, start=1):
            doc.add_paragraph(f"{i}. {q}")
        doc.add_paragraph("")
    doc.save(path)


if __name__ == "__main__":
    students = os.path.join(DATA, "students.xlsx")
    tickets = os.path.join(DATA, "tickets.docx")
    write_students(students)
    write_tickets(tickets)
    shutil.copy2(students, os.path.join(ROOT, "students.xlsx"))
    shutil.copy2(tickets, os.path.join(ROOT, "tickets.docx"))
    print("Созданы data/students.xlsx, data/tickets.docx и копии в корне проекта.")
