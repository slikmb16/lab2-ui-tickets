"""
UI-тесты веб-интерфейса (Selenium).

Прогоняют сценарии приёмки кликами: группы, пустая группа,
активность кнопки, окно билета, ESC, идемпотентность.
"""
import os
import shutil
import sys
import tempfile
import threading
import time
import unittest
from wsgiref.simple_server import make_server

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx import Document
from openpyxl import Workbook

from app.web import create_app

try:
    from selenium import webdriver
    from selenium.webdriver.chrome.options import Options
    from selenium.webdriver.common.by import By
    from selenium.webdriver.common.keys import Keys
    from selenium.webdriver.support import expected_conditions as EC
    from selenium.webdriver.support.ui import Select, WebDriverWait
except ImportError:  # pragma: no cover
    webdriver = None


def _make_students(path):
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("ИС-21")
    ws.append(["Фамилия", "Имя"])
    ws.append(["Иванов", "Иван"])
    ws.append(["Петрова", "Анна"])
    ws2 = wb.create_sheet("ПИ-22")
    ws2.append(["Фамилия", "Имя"])
    ws2.append(["Сидоров", "Семён"])
    ws3 = wb.create_sheet("АРХ-20")
    ws3.append(["Фамилия", "Имя"])
    wb.save(path)


def _make_tickets(path):
    doc = Document()
    doc.add_paragraph("Билет 1")
    doc.add_paragraph("1. Что такое алгоритм?")
    doc.add_paragraph("2. Чем стек отличается от очереди?")
    doc.add_paragraph("3. Что такое рекурсия?")
    doc.add_paragraph("")
    doc.add_paragraph("Билет 2")
    doc.add_paragraph("1. Что такое О-большое?")
    doc.add_paragraph("2. Как работает бинарный поиск?")
    doc.add_paragraph("3. Что такое хеш-таблица?")
    doc.add_paragraph("")
    doc.add_paragraph("Билет 4")
    doc.add_paragraph("1. Битый билет с двумя вопросами")
    doc.add_paragraph("2. Этот билет должен быть пропущен")
    doc.save(path)


@unittest.skipIf(webdriver is None, "selenium не установлен")
class TestWebUISelenium(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmpdir = tempfile.mkdtemp()
        _make_students(os.path.join(cls.tmpdir, "students.xlsx"))
        _make_tickets(os.path.join(cls.tmpdir, "tickets.docx"))
        src_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        for name in ("templates", "static"):
            shutil.copytree(os.path.join(src_root, name), os.path.join(cls.tmpdir, name))

        app = create_app(cls.tmpdir)
        cls.httpd = make_server("127.0.0.1", 0, app)
        cls.port = cls.httpd.server_port
        cls.thread = threading.Thread(target=cls.httpd.serve_forever, daemon=True)
        cls.thread.start()

        options = Options()
        options.add_argument("--headless=new")
        options.add_argument("--disable-gpu")
        options.add_argument("--window-size=1400,900")
        options.add_argument("--no-sandbox")
        try:
            cls.driver = webdriver.Chrome(options=options)
        except Exception as exc:  # pragma: no cover
            cls.httpd.shutdown()
            raise unittest.SkipTest(f"ChromeDriver недоступен: {exc}") from exc

        cls.wait = WebDriverWait(cls.driver, 8)
        cls.base = f"http://127.0.0.1:{cls.port}"

    @classmethod
    def tearDownClass(cls):
        if getattr(cls, "driver", None):
            cls.driver.quit()
        if getattr(cls, "httpd", None):
            cls.httpd.shutdown()
        shutil.rmtree(cls.tmpdir, ignore_errors=True)

    def setUp(self):
        self.driver.get(self.base)
        self.wait.until(EC.presence_of_element_located((By.CSS_SELECTOR, '[data-testid="group-select"]')))
        self.wait.until(lambda d: d.find_element(By.ID, "stat-groups").text not in ("", "—"))

    def test_groups_loaded_and_generate_disabled_until_student(self):
        group = Select(self.driver.find_element(By.CSS_SELECTOR, '[data-testid="group-select"]'))
        names = [o.text for o in group.options]
        self.assertIn("ИС-21", names)
        self.assertIn("АРХ-20", names)
        btn = self.driver.find_element(By.CSS_SELECTOR, '[data-testid="generate-btn"]')
        self.assertFalse(btn.is_enabled())
        group.select_by_visible_text("ИС-21")
        time.sleep(0.3)
        self.assertFalse(btn.is_enabled())
        Select(self.driver.find_element(By.CSS_SELECTOR, '[data-testid="student-select"]')).select_by_visible_text(
            "Иванов Иван"
        )
        self.assertTrue(btn.is_enabled())

    def test_empty_group_disables_student_and_generate(self):
        Select(self.driver.find_element(By.CSS_SELECTOR, '[data-testid="group-select"]')).select_by_visible_text(
            "АРХ-20"
        )
        time.sleep(0.3)
        student = self.driver.find_element(By.CSS_SELECTOR, '[data-testid="student-select"]')
        self.assertFalse(student.is_enabled())
        btn = self.driver.find_element(By.CSS_SELECTOR, '[data-testid="generate-btn"]')
        self.assertFalse(btn.is_enabled())
        hint = self.driver.find_element(By.ID, "hint").text
        self.assertIn("нет студентов", hint)

    def test_generate_shows_ticket_and_three_questions_esc_resets(self):
        Select(self.driver.find_element(By.CSS_SELECTOR, '[data-testid="group-select"]')).select_by_visible_text(
            "ИС-21"
        )
        time.sleep(0.3)
        Select(self.driver.find_element(By.CSS_SELECTOR, '[data-testid="student-select"]')).select_by_visible_text(
            "Петрова Анна"
        )
        self.driver.find_element(By.CSS_SELECTOR, '[data-testid="generate-btn"]').click()
        modal = self.wait.until(EC.visibility_of_element_located((By.CSS_SELECTOR, '[data-testid="result-modal"]')))
        title = self.driver.find_element(By.CSS_SELECTOR, '[data-testid="ticket-title"]').text
        self.assertRegex(title, r"Ваш билет № \d+")
        questions = self.driver.find_elements(By.CSS_SELECTOR, ".question")
        self.assertEqual(len(questions), 3)
        self.driver.find_element(By.TAG_NAME, "body").send_keys(Keys.ESCAPE)
        self.wait.until(EC.invisibility_of_element(modal))
        btn = self.driver.find_element(By.CSS_SELECTOR, '[data-testid="generate-btn"]')
        self.assertFalse(btn.is_enabled())
        self.assertTrue(self.driver.find_element(By.CSS_SELECTOR, '[data-testid="group-select"]').is_displayed())

    def test_same_student_gets_same_ticket_number(self):
        def generate_once():
            Select(self.driver.find_element(By.CSS_SELECTOR, '[data-testid="group-select"]')).select_by_visible_text(
                "ПИ-22"
            )
            time.sleep(0.3)
            Select(self.driver.find_element(By.CSS_SELECTOR, '[data-testid="student-select"]')).select_by_visible_text(
                "Сидоров Семён"
            )
            self.driver.find_element(By.CSS_SELECTOR, '[data-testid="generate-btn"]').click()
            self.wait.until(EC.visibility_of_element_located((By.CSS_SELECTOR, '[data-testid="result-modal"]')))
            title = self.driver.find_element(By.CSS_SELECTOR, '[data-testid="ticket-title"]').text
            self.driver.find_element(By.TAG_NAME, "body").send_keys(Keys.ESCAPE)
            self.wait.until(EC.invisibility_of_element_located((By.CSS_SELECTOR, '[data-testid="result-modal"]')))
            return title

        first = generate_once()
        second = generate_once()
        self.assertEqual(first, second)
        self.assertRegex(first, r"Ваш билет № \d+")


if __name__ == "__main__":
    unittest.main(verbosity=2)
