"""
Unit-тесты:
- парсинг билета из tickets.docx (включая устойчивость к битым билетам);
- парсинг студентов из students.xlsx (включая пустую группу);
- поиск студента по группе / идемпотентность выбора билета;
- запись и дозапись results.xlsx (повторы, порядок, неизменность старых строк).
"""
import os
import shutil
import sys
import tempfile
import unittest
from unittest.mock import patch

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx import Document
from openpyxl import Workbook, load_workbook

from app import logic, parsing, storage
from app.web import create_app


class TestTicketParsing(unittest.TestCase):
    def _make_docx(self, path, ticket_blocks):
        doc = Document()
        for number, questions in ticket_blocks:
            doc.add_paragraph(f"Билет {number}")
            for i, q in enumerate(questions, start=1):
                doc.add_paragraph(f"{i}. {q}")
            doc.add_paragraph("")
        doc.save(path)

    def test_parses_valid_tickets(self):
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "tickets.docx")
            self._make_docx(path, [
                (1, ["Вопрос 1а", "Вопрос 1б", "Вопрос 1в"]),
                (2, ["Вопрос 2а", "Вопрос 2б", "Вопрос 2в"]),
            ])
            tickets = parsing.load_tickets(path)
            self.assertEqual(set(tickets.keys()), {1, 2})
            self.assertEqual(tickets[1], ["Вопрос 1а", "Вопрос 1б", "Вопрос 1в"])
            self.assertEqual(len(tickets[2]), 3)

    def test_skips_ticket_with_fewer_than_3_questions(self):
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "tickets.docx")
            self._make_docx(path, [
                (1, ["Только один вопрос"]),
                (2, ["Вопрос 2а", "Вопрос 2б", "Вопрос 2в"]),
            ])
            tickets = parsing.load_tickets(path)
            self.assertNotIn(1, tickets)
            self.assertIn(2, tickets)

    def test_unrecognized_header_does_not_crash_parsing(self):
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "tickets.docx")
            doc = Document()
            doc.add_paragraph("Это не заголовок билета")
            doc.add_paragraph("1. Какой-то текст")
            doc.add_paragraph("Билет 7")
            doc.add_paragraph("1. Вопрос а")
            doc.add_paragraph("2. Вопрос б")
            doc.add_paragraph("3. Вопрос в")
            doc.save(path)

            tickets = parsing.load_tickets(path)
            self.assertEqual(set(tickets.keys()), {7})

    def test_extra_questions_truncated_to_three(self):
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "tickets.docx")
            self._make_docx(path, [
                (1, ["Q1", "Q2", "Q3", "Q4"]),
            ])
            tickets = parsing.load_tickets(path)
            self.assertEqual(len(tickets[1]), 3)


class TestStudentParsing(unittest.TestCase):
    def _make_xlsx(self, path, sheets):
        wb = Workbook()
        wb.remove(wb.active)
        for name, rows in sheets.items():
            ws = wb.create_sheet(name)
            ws.append(["Фамилия", "Имя"])
            for row in rows:
                ws.append(row)
        wb.save(path)

    def test_parses_students_per_group(self):
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "students.xlsx")
            self._make_xlsx(path, {
                "Группа1": [["Иванов", "Иван"], ["Петров", "Пётр"]],
                "Группа2": [["Сидоров", "Семён"]],
            })
            groups = parsing.load_students(path)
            self.assertEqual(len(groups["Группа1"]), 2)
            self.assertEqual(groups["Группа1"][0].full_name, "Иванов Иван")
            self.assertEqual(len(groups["Группа2"]), 1)

    def test_empty_group_returns_empty_list_no_exception(self):
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "students.xlsx")
            self._make_xlsx(path, {"ПустаяГруппа": []})
            groups = parsing.load_students(path)
            self.assertEqual(groups["ПустаяГруппа"], [])


class TestAssignmentIdempotency(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()
        self.results_path = os.path.join(self.tmpdir, "results.xlsx")
        self.tickets = {1: ["a", "b", "c"], 2: ["d", "e", "f"], 3: ["g", "h", "i"]}

    def tearDown(self):
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def test_new_student_gets_random_ticket_marked_not_repeat(self):
        ticket_number, repeat = logic.assign_ticket(
            self.results_path, "ГрА", "Иванов", "Иван", self.tickets
        )
        self.assertIn(ticket_number, self.tickets)
        self.assertFalse(repeat)

    def test_repeated_student_gets_same_ticket_marked_repeat(self):
        first_ticket, first_repeat = logic.assign_ticket(
            self.results_path, "ГрА", "Иванов", "Иван", self.tickets
        )
        storage.append_result(
            self.results_path, "ГрА", "Иванов", "Иван", first_ticket, first_repeat
        )

        second_ticket, second_repeat = logic.assign_ticket(
            self.results_path, "ГрА", "Иванов", "Иван", self.tickets
        )
        self.assertEqual(second_ticket, first_ticket)
        self.assertTrue(second_repeat)

    def test_first_recorded_ticket_wins_even_after_many_repeats(self):
        # Первая генерация
        t1, r1 = logic.assign_ticket(self.results_path, "ГрА", "Петров", "Пётр", self.tickets)
        storage.append_result(self.results_path, "ГрА", "Петров", "Пётр", t1, r1)

        # Несколько повторных обращений подряд
        for _ in range(3):
            tn, rn = logic.assign_ticket(self.results_path, "ГрА", "Петров", "Пётр", self.tickets)
            storage.append_result(self.results_path, "ГрА", "Петров", "Пётр", tn, rn)
            self.assertEqual(tn, t1)
            self.assertTrue(rn)

        # В журнале должно быть 4 строки, все с одним и тем же номером билета
        wb = load_workbook(self.results_path)
        ws = wb.active
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(rows), 4)
        self.assertTrue(all(row[3] == t1 for row in rows))
        self.assertEqual(rows[0][5], "нет")
        self.assertTrue(all(row[5] == "да" for row in rows[1:]))

    def test_old_rows_are_not_modified_on_append(self):
        storage.append_result(self.results_path, "ГрА", "Иванов", "Иван", 1, False)
        wb = load_workbook(self.results_path)
        first_row_before = list(wb.active.iter_rows(min_row=2, max_row=2, values_only=True))[0]

        storage.append_result(self.results_path, "ГрБ", "Смирнов", "Олег", 2, False)
        wb2 = load_workbook(self.results_path)
        rows_after = list(wb2.active.iter_rows(min_row=2, values_only=True))

        self.assertEqual(rows_after[0][:4], first_row_before[:4])
        self.assertEqual(len(rows_after), 2)

    def test_different_groups_same_name_are_independent(self):
        t1, r1 = logic.assign_ticket(self.results_path, "ГрА", "Иванов", "Иван", self.tickets)
        storage.append_result(self.results_path, "ГрА", "Иванов", "Иван", t1, r1)

        # Тот же "Иванов Иван", но другая группа — считается новым студентом
        t2, r2 = logic.assign_ticket(self.results_path, "ГрБ", "Иванов", "Иван", self.tickets)
        self.assertFalse(r2)


class TestResultsFileCreation(unittest.TestCase):
    def test_creates_file_with_correct_headers_if_missing(self):
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "results.xlsx")
            self.assertFalse(os.path.exists(path))
            storage.ensure_results_file(path)
            self.assertTrue(os.path.exists(path))

            wb = load_workbook(path)
            header = list(wb.active.iter_rows(min_row=1, max_row=1, values_only=True))[0]
            self.assertEqual(list(header), storage.HEADERS)

    def test_does_not_overwrite_existing_file(self):
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "results.xlsx")
            storage.append_result(path, "ГрА", "Иванов", "Иван", 1, False)
            storage.ensure_results_file(path)  # не должно ничего сломать/перезаписать

            wb = load_workbook(path)
            rows = list(wb.active.iter_rows(min_row=2, values_only=True))
            self.assertEqual(len(rows), 1)


class TestMandatoryRequirements(unittest.TestCase):
    def test_missing_input_files_return_clear_fatal_status(self):
        """Без students.xlsx и tickets.docx приложение остаётся доступным."""
        with tempfile.TemporaryDirectory() as d:
            # Даже наличие старых копий в data/ не должно скрывать ошибку:
            # по лабораторной входные файлы ожидаются в папке проекта.
            os.mkdir(os.path.join(d, "data"))
            open(os.path.join(d, "data", "students.xlsx"), "wb").close()
            open(os.path.join(d, "data", "tickets.docx"), "wb").close()
            response = create_app(d).test_client().get("/api/status")
            data = response.get_json()
            self.assertEqual(response.status_code, 200)
            self.assertIn("students.xlsx", data["fatal"])
            self.assertIn("tickets.docx", data["fatal"])

    def test_locked_results_file_raises_friendly_error(self):
        """Ошибка доступа к Excel превращается в контролируемую ошибку приложения."""
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "results.xlsx")
            storage.ensure_results_file(path)
            with patch("app.storage.load_workbook", side_effect=PermissionError("locked")):
                with self.assertRaises(storage.ResultsFileLocked):
                    storage.append_result(path, "ГрА", "Иванов", "Иван", 1, False, retries=1)


if __name__ == "__main__":
    unittest.main(verbosity=2)
